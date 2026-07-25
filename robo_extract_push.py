import openpyxl, json
from openpyxl.utils import get_column_letter as L
from collections import defaultdict, OrderedDict
import sys
wb=openpyxl.load_workbook(sys.argv[1] if len(sys.argv)>1 else "planilha.xlsx",data_only=True)

MES_ORD=["Janeiro","Fevereiro","Março","Abril","Maio","Junho","Julho","Agosto","Setembro","Outubro","Novembro","Dezembro"]
MES_ABR={"Janeiro":"Jan","Fevereiro":"Fev","Março":"Mar","Abril":"Abr","Maio":"Mai","Junho":"Jun","Julho":"Jul","Agosto":"Ago","Setembro":"Set","Outubro":"Out","Novembro":"Nov","Dezembro":"Dez"}

# ---------- BASE (realizado) ----------
b=wb["Base"]; H={b.cell(2,c).value:c for c in range(1,b.max_column+1) if b.cell(2,c).value}
def bc(r,name): 
    c=H.get(name); return b.cell(r,c).value if c else None
prodYTD=defaultdict(lambda:{"real":0.0,"qtd":0})
prodMensal=defaultdict(lambda:[0.0]*12)
canalYTD=defaultdict(lambda:{"real":0.0,"qtd":0})
canalMensal=defaultdict(lambda:[0.0]*12)
gerYTD=defaultdict(lambda:{"real":0.0,"qtd":0})
corrYTD=defaultdict(lambda:{"real":0.0,"qtd":0})
mensalTotal=[0.0]*12
totReal=0.0; totQtd=0
# nomes canônicos de empreendimento (da aba KPIs) para normalizar grafias divergentes na Base
_kk=wb["KPIs"]
_canon=[str(_kk.cell(r,1).value).strip() for r in range(3,14) if _kk.cell(r,1).value]
_normMap={n.upper():n for n in _canon}
def canon(p):
    s=str(p).strip(); return _normMap.get(s.upper(), s)
for r in range(3,b.max_row+1):
    p=bc(r,"Produto"); v=bc(r,"Valor")
    if p is None and v is None: continue
    try: v=float(v or 0)
    except: v=0
    mesn=bc(r,"MÊS2") or bc(r,"Mês")
    mi=None
    if isinstance(mesn,str) and mesn.strip() in MES_ORD: mi=MES_ORD.index(mesn.strip())
    else:
        d=bc(r,"Mês")
        if hasattr(d,"month"): mi=d.month-1
    p=canon(p)
    canal=str(bc(r,"Canal") or "—").strip()
    ger=str(bc(r,"Gerente") or "—").strip()
    corr=bc(r,"Corretor"); corr=str(corr).strip() if corr and str(corr).strip() else None
    prodYTD[p]["real"]+=v; prodYTD[p]["qtd"]+=1
    canalYTD[canal]["real"]+=v; canalYTD[canal]["qtd"]+=1
    gerYTD[ger]["real"]+=v; gerYTD[ger]["qtd"]+=1
    if corr: corrYTD[corr]["real"]+=v; corrYTD[corr]["qtd"]+=1
    totReal+=v; totQtd+=1
    if mi is not None:
        prodMensal[p][mi]+=v; canalMensal[canal][mi]+=v; mensalTotal[mi]+=v

# ---------- METAS (das abas de plano) ----------
k=wb["KPIs"]; kc=lambda r,c:k.cell(r,c).value
# monthly total meta + curva (L14:Q25)
meses=[]
for i,r in enumerate(range(14,26)):
    meses.append(dict(mes=MES_ABR[MES_ORD[i]], meta=kc(r,13), real=round(mensalTotal[i],2),
                      metaAcum=kc(r,15)))
# real acum from base
acc=0
for m in meses: acc+=m["real"]; m["realAcum"]=round(acc,2)
metaAnoTotal=kc(26,13)
# product meta YTD (KPIs rows 3-13)
prodMetaYTD={}
for r in range(3,14):
    nm=kc(r,1)
    if nm and kc(r,2) is not None: prodMetaYTD[str(nm).strip()]=kc(r,2)
prodVSO={str(kc(r,1)).strip():kc(r,10) for r in range(3,14) if kc(r,1)}
metaYTD=kc(14,2); 
# canal meta YTD (rows 39-42)
canalMeta={}
for r in range(39,43):
    nm=kc(r,1)
    if nm: canalMeta[str(nm).strip()]=kc(r,2) if isinstance(kc(r,2),(int,float)) else None
canalShare={str(kc(r,1)).strip():kc(r,5) for r in range(39,43) if kc(r,1)}
# campanha seleção (KPIs rows 47-51)
campanha=[]
for r in range(48,52):
    nm=kc(r,1)
    if nm: campanha.append(dict(produto=str(nm).strip(),periodo=str(kc(r,2)),vgv=kc(r,3),un=kc(r,4),unAno=kc(r,5),peso=kc(r,6)))

# gerente meta YTD (POR GERENTE: meta cols B,F,J,... = 2+4*m for months 0..6)
pg=wb["POR GERENTE"]
gerMeta=defaultdict(float); gerRealTab=defaultdict(float)
for r in range(6,pg.max_row+1):
    nm=pg.cell(r,1).value
    if not nm or str(nm).strip() in ("Total","TOTAL",""): 
        if str(nm).strip() in ("Total","TOTAL"): break
        continue
    nm=str(nm).strip()
    for mo in range(7):  # Jan-Jul
        mv=pg.cell(r,2+4*mo).value
        rv=pg.cell(r,3+4*mo).value
        if isinstance(mv,(int,float)): gerMeta[nm]+=mv
        if isinstance(rv,(int,float)): gerRealTab[nm]+=rv

# ---------- POR SEMANA (matriz semanal GERAL) ----------
sw=wb["POR SEMANA"]
# week blocks start col2 step4; month from row2 carry; range from row3
weeks=[]
curmonth=""
import re
c=2
while c<=sw.max_column-1:
    mo=sw.cell(2,c).value
    if mo: curmonth=str(mo).strip()
    rng=sw.cell(3,c).value
    # manter apenas blocos semanais reais (range "NN a NN"); descartar resumo "Meta mês"
    if rng and re.match(r'^\d+\s*a\s*\d+$',str(rng).strip()):
        weeks.append(dict(col=c,mes=curmonth,rng=str(rng).strip()))
    c+=4
# 4 blocos por canal (mesmas colunas de semana, faixas de linha diferentes)
SEM_BLOCOS={"Geral":(5,15),"Salão":(22,32),"Online":(39,49),"Parcerias":(56,66)}
def sem_block(r0,r1):
    out=[]
    for r in range(r0,r1+1):
        nm=sw.cell(r,1).value
        if not nm or str(nm).strip() in ("TOTAL",""): continue
        nm=canon(nm); arr=[]
        for w in weeks:
            cc=w["col"]
            meta=sw.cell(r,cc).value or 0; real=sw.cell(r,cc+1).value or 0; qtd=sw.cell(r,cc+3).value or 0
            arr.append([round(meta,2) if isinstance(meta,(int,float)) else 0,
                        round(real,2) if isinstance(real,(int,float)) else 0,
                        qtd if isinstance(qtd,(int,float)) else 0])
        out.append(dict(nome=nm, w=arr))
    return out
semCanais={c:sem_block(*rg) for c,rg in SEM_BLOCOS.items()}
semProds=semCanais["Geral"]
semLabels=[dict(mes=MES_ABR.get(w["mes"],w["mes"][:3]),rng=w["rng"]) for w in weeks]

# ---------- montar produtos com gap ----------
produtos=[]
for nm,mYTD in prodMetaYTD.items():
    real=round(prodYTD.get(nm,{}).get("real",0),2)
    qtd=prodYTD.get(nm,{}).get("qtd",0)
    gap=round(mYTD-real,2)
    produtos.append(dict(nome=nm,metaYTD=round(mYTD,2),real=real,qtd=qtd,gap=gap,
                         ating=round(real/mYTD,4) if mYTD else 0,
                         vso=prodVSO.get(nm),
                         mensal=[round(x,2) for x in prodMensal.get(nm,[0]*12)]))
# ordena por gap desc? user: primeiro menor gap, ultimo pior gap -> sort by gap ascending
produtos.sort(key=lambda x:x["gap"])

canais=[]
for nm in ["Salão","Online","Parcerias","Interna"]:
    real=round(canalYTD.get(nm,{}).get("real",0),2)
    qtd=canalYTD.get(nm,{}).get("qtd",0)
    mt=canalMeta.get(nm)
    canais.append(dict(canal=nm,meta=mt,real=real,qtd=qtd,
                       ating=round(real/mt,4) if isinstance(mt,(int,float)) and mt else None,
                       share=round(real/totReal,4),
                       mensal=[round(x,2) for x in canalMensal.get(nm,[0]*12)]))

gerentes=[]
for nm,d in gerYTD.items():
    if nm in ("—","Interna"): 
        pass
    real=round(d["real"],2); qtd=d["qtd"]
    mt=round(gerMeta.get(nm,0),2)
    gerentes.append(dict(nome=nm,real=real,qtd=qtd,meta=mt,
                         ating=round(real/mt,4) if mt else None,
                         ticket=round(real/qtd,2) if qtd else 0))
gerentes.sort(key=lambda x:-x["real"])

corretores=[]
for nm,d in corrYTD.items():
    corretores.append(dict(nome=nm,real=round(d["real"],2),qtd=d["qtd"],
                           ticket=round(d["real"]/d["qtd"],2) if d["qtd"] else 0))
corretores.sort(key=lambda x:-x["real"])

DADOS=dict(
  atualizado="20/07/2026", empresa="Planik Empreendimentos Imobiliários",
  periodo="Ano 2026 · corte YTD",
  metaAnoTotal=round(metaAnoTotal,2), metaYTD=round(metaYTD,2),
  realYTD=round(totReal,2), qtdYTD=totQtd,
  atingYTD=round(totReal/metaYTD,4),
  meses=meses, produtos=produtos, canais=canais,
  gerentes=gerentes, corretores=corretores, campanha=campanha,
  semana=dict(labels=semLabels, produtos=semProds, canais=semCanais),
  corrCobertura=round(sum(d["qtd"] for d in corrYTD.values())/totQtd,3)
)
json.dump(DADOS,open("dados_full.json","w",encoding="utf-8"),ensure_ascii=False,default=str)

# ---------- VERIFICAÇÃO ----------
print("VERIFICAÇÃO")
print(f"  Real YTD Base: {totReal/1e6:.2f} mi | KPIs esperado: 231.94 mi | Qtd: {totQtd}")
print(f"  Meta YTD: {metaYTD/1e6:.2f} mi | Meta ano total: {metaAnoTotal/1e6:.2f} mi | Ating: {totReal/metaYTD*100:.1f}%")
print(f"  Semanas extraídas: {len(weeks)} | Produtos semana: {len(semProds)}")
print(f"  Gerentes: {len(gerentes)} | Corretores nomeados: {len(corretores)} (cobertura {DADOS['corrCobertura']*100:.0f}%)")
print("  Canal real (Base):", {c['canal']:round(c['real']/1e6,1) for c in canais})
print("  Top 3 gap (melhor):", [(p['nome'],round(p['gap']/1e6,1)) for p in produtos[:3]])
print("  Pior gap:", [(p['nome'],round(p['gap']/1e6,1)) for p in produtos[-3:]])
print("  Real por gerente (top5):", [(g['nome'],round(g['real']/1e6,1),g['qtd']) for g in gerentes[:5]])

# ================= V4 ADICIONAIS =================
import datetime
# meta total real = Meta Acum Dezembro (O25) = 479,58 mi (forecast redistribuído não soma)
metaAnoReal=kc(25,15)

# gerente x canal (da Base)
gerCanal=defaultdict(lambda:defaultdict(lambda:{"real":0.0,"qtd":0}))
transacoes=[]
for r in range(3,b.max_row+1):
    p=bc(r,"Produto"); v=bc(r,"Valor")
    if p is None and v is None: continue
    try: v=float(v or 0)
    except: v=0
    ger=str(bc(r,"Gerente") or "—").strip()
    canal=str(bc(r,"Canal") or "—").strip()
    gerCanal[ger][canal]["real"]+=v; gerCanal[ger][canal]["qtd"]+=1
    d=bc(r,"Mês")
    diso=None
    if hasattr(d,"isoformat"): diso=d.date().isoformat() if hasattr(d,"date") else d.isoformat()
    uni=bc(r,"Unidade"); cor=bc(r,"Corretor"); desc=bc(r,"Desconto")
    transacoes.append({"p":canon(p),"u":str(uni).strip() if uni else "—","d":diso,
                       "cor":str(cor).strip() if cor and str(cor).strip() else "—",
                       "ger":ger,"c":canal,"v":round(v,2),
                       "desc":round(float(desc),4) if isinstance(desc,(int,float)) else None})

ACT_PLANIK=["Dorival","Ramon","Thais"]
ACT_PARC=["Edmond","Lorrane","Pedro","Cecilia","Kalu"]
def gtot(nm,canais):
    d=gerCanal.get(nm,{}); return round(sum(d[c]["real"] for c in canais if c in d),2), sum(d[c]["qtd"] for c in canais if c in d)
gerAtivos={"planik":[],"parc":[],"geral":[]}
for nm in ACT_PLANIK:
    rl,q=gtot(nm,["Salão","Online"]); mt=round(gerMeta.get(nm,0),2)
    gerAtivos["planik"].append(dict(nome=nm,real=rl,qtd=q,meta=mt,ating=round(rl/mt,4) if mt else None,ticket=round(rl/q,2) if q else 0))
for nm in ACT_PARC:
    rl,q=gtot(nm,["Parcerias"]); mt=round(gerMeta.get(nm,0),2)
    gerAtivos["parc"].append(dict(nome=nm,real=rl,qtd=q,meta=mt,ating=round(rl/mt,4) if mt else None,ticket=round(rl/q,2) if q else 0))
# geral = todos ativos, VGV total
for nm in ACT_PLANIK+ACT_PARC:
    d=gerCanal.get(nm,{}); rl=round(sum(x["real"] for x in d.values()),2); q=sum(x["qtd"] for x in d.values())
    mt=round(gerMeta.get(nm,0),2)
    time="Planik Vendas" if nm in ACT_PLANIK else "Parcerias"
    gerAtivos["geral"].append(dict(nome=nm,real=rl,qtd=q,meta=mt,ating=round(rl/mt,4) if mt else None,ticket=round(rl/q,2) if q else 0,time=time))
gerAtivos["geral"].sort(key=lambda x:-x["real"])
gerAtivos["planik"].sort(key=lambda x:-x["real"])
gerAtivos["parc"].sort(key=lambda x:-x["real"])

# dados por gerente x canal (para recalcular ativos dinamicamente na engrenagem)
gerCanalData={}
for nm,d in gerCanal.items():
    canais={c:{"real":round(x["real"],2),"qtd":x["qtd"]} for c,x in d.items()}
    tr=round(sum(x["real"] for x in d.values()),2); tq=sum(x["qtd"] for x in d.values())
    gerCanalData[nm]=dict(canais=canais,real=tr,qtd=tq,meta=round(gerMeta.get(nm,0),2))

# ---- DESEMPENHO CORRETORES ----
dc=wb["DESEMPENHO CORRETORES"]
corrKPIs=dict(vgv=dc.cell(6,1).value,vendas=dc.cell(6,3).value,ticket=dc.cell(6,5).value,
              corretores=dc.cell(6,7).value,pctAtivos=dc.cell(6,9).value,vgvPorCorr=dc.cell(6,11).value)
corrEquipes=[]
for r in range(9,12):
    corrEquipes.append(dict(gerente=dc.cell(r,1).value,corretores=dc.cell(r,2).value,vendeu3=dc.cell(r,3).value,
        pctVendeu=dc.cell(r,4).value,vendas=dc.cell(r,5).value,vgv=dc.cell(r,6).value,part=dc.cell(r,7).value,
        ticket=dc.cell(r,8).value,vgvCorr=dc.cell(r,9).value,pctSalao=dc.cell(r,10).value,pctOnline=dc.cell(r,11).value,conc=dc.cell(r,12).value))
corrRanking=[]
for r in range(16,89):
    nome=dc.cell(r,2).value
    if not nome: continue
    vgv=dc.cell(r,6).value or 0
    uv=dc.cell(r,9).value
    corrRanking.append(dict(nome=str(nome).strip(),equipe=str(dc.cell(r,3).value or "").strip(),
        canal=str(dc.cell(r,4).value or "").strip(),qtd=dc.cell(r,5).value or 0,vgv=round(float(vgv),2),
        pctTotal=dc.cell(r,7).value,ticket=dc.cell(r,8).value,
        ultima=(uv.date().isoformat() if hasattr(uv,"date") else str(uv) if uv else None),
        recencia=dc.cell(r,10).value,naoVendeu=(dc.cell(r,11).value=="Sim")))

DADOS["metaAnoReal"]=round(metaAnoReal,2)
DADOS["gerAtivos"]=gerAtivos
DADOS["gerCanalData"]=gerCanalData
DADOS["ativosSeed"]={"planik":ACT_PLANIK,"parc":ACT_PARC}
DADOS["gerCanalNota"]="Ativos: Planik Vendas (Dorival, Ramon, Thais) · Parcerias (Edmond, Lorrane, Pedro, Cecilia, Kalu)"
DADOS["transacoes"]=transacoes
# Data de atualização: puxa da célula "Data atualização: dd/mm/aaaa" no topo da Base (A1).
# Se não achar, usa a venda mais recente como reserva.
import re as _re
from datetime import datetime as _dt
_atual=None
_a1=b.cell(1,1).value
if isinstance(_a1,_dt):
    _atual=_a1.strftime("%d/%m/%Y")
elif _a1 is not None:
    _m=_re.search(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})",str(_a1))
    if _m:
        dd,mm,yy=_m.groups(); yy=("20"+yy) if len(yy)==2 else yy
        try: _atual=_dt(int(yy),int(mm),int(dd)).strftime("%d/%m/%Y")
        except Exception as _e: print("aviso data A1:",_e)
if not _atual:
    _dts=[t["d"] for t in transacoes if t.get("d")]
    if _dts:
        try: _atual=_dt.strptime(max(_dts),"%Y-%m-%d").strftime("%d/%m/%Y")
        except Exception as _e: print("aviso data venda:",_e)
if _atual: DADOS["atualizado"]=_atual
print("  Data atualização (A1):",repr(_a1),"->",DADOS["atualizado"])
DADOS["corr"]=dict(kpis=corrKPIs,equipes=corrEquipes,ranking=corrRanking)
json.dump(DADOS,open("dados_full.json","w",encoding="utf-8"),ensure_ascii=False,default=str)
print("\nV4:")
print("  Meta ano REAL (O25):",round(metaAnoReal/1e6,2),"mi (antes usava 592.47)")
print("  Gerentes ativos geral:",[(g['nome'],round(g['real']/1e6,1),g['time']) for g in gerAtivos['geral']])
print("  Transações:",len(transacoes),"| datas ex:",transacoes[0]['d'],transacoes[-1]['d'])
print("  Corr KPIs: VGV",round(corrKPIs['vgv']/1e6,1),"mi | vendas",corrKPIs['vendas'],"| corretores",corrKPIs['corretores'])
print("  Corr ranking linhas:",len(corrRanking),"| com venda:",sum(1 for c in corrRanking if c['qtd']>0))
print("  Canais no ranking:",set(c['canal'] for c in corrRanking))

# ================= V5 ADICIONAIS =================
# Síntese "vida do empreendimento" (Por Canal RESUMIDO r30-40)
pcr=wb["Por Canal RESUMIDO"]
sintese=[]
for r in range(30,41):
    nm=pcr.cell(r,1).value
    if not nm or str(nm).strip().startswith("*"): continue
    g=lambda c: pcr.cell(r,c).value
    sintese.append(dict(produto=canon(nm), m2=g(2), vgvTotal=g(3), vendidoVida=g(4), disponivel=g(5),
                        uniTotal=g(6), uniVendidas=g(7), permutas=g(8), livres=g(9), pctVgv=g(10), pctQtd=g(11)))

# VSO mês a mês (VSOs 2026) — total + por produto
vsows=wb["VSOs 2026"]
MESES_UP=["JANEIRO","FEVEREIRO","MARÇO","ABRIL","MAIO","JUNHO","JULHO","AGOSTO","SETEMBRO","OUTUBRO","NOVEMBRO","DEZEMBRO"]
locs={}
for r in range(1,vsows.max_row+1):
    for c in range(1,vsows.max_column+1):
        v=vsows.cell(r,c).value
        if isinstance(v,str) and v.strip().upper() in MESES_UP: locs[v.strip().upper()]=(r,c)
vsoMensal=[]; vsoProd=defaultdict(lambda:[None]*12)
for mi,mu in enumerate(MESES_UP):
    if mu not in locs: vsoMensal.append(dict(mes=MES_ABR[MES_ORD[mi]],vso=None,vend=None,disp=None)); continue
    r0,c0=locs[mu]; tot=None
    for r in range(r0+2,r0+20):
        nm=vsows.cell(r,c0).value
        if nm is None: break
        disp=vsows.cell(r,c0+1).value; vend=vsows.cell(r,c0+2).value; vso=vsows.cell(r,c0+3).value
        if str(nm).strip().lower()=="total": tot=(disp,vend,vso); break
        if isinstance(vso,(int,float)): vsoProd[canon(nm)][mi]=round(vso,4)
    numf=lambda x: round(x,4) if isinstance(x,(int,float)) else None
    if tot: vsoMensal.append(dict(mes=MES_ABR[MES_ORD[mi]],vso=numf(tot[2]),
                                  vend=numf(tot[1]),disp=numf(tot[0])))
    else: vsoMensal.append(dict(mes=MES_ABR[MES_ORD[mi]],vso=None,vend=None,disp=None))

DADOS["sintese"]=sintese
DADOS["vsoMensal"]=vsoMensal
DADOS["vsoProd"]={k:v for k,v in vsoProd.items()}
json.dump(DADOS,open("dados_full.json","w",encoding="utf-8"),ensure_ascii=False,default=str)
print("\nV5:")
print("  Síntese produtos:",len(sintese),"| ex:",sintese[0]['produto'],"m2",sintese[0]['m2'],"VGV total",round(sintese[0]['vgvTotal']/1e6,1),"vendido vida%",round(sintese[0]['pctVgv']*100,1))
print("  VSO mensal:",[(v['mes'],round(v['vso']*100,1) if v['vso'] else None) for v in vsoMensal])
print("  VSO por produto (ex NIK FREI CANECA):",vsoProd.get("NIK FREI CANECA"))

# ================= V6: POR PRODUTO YTD (por canal, mês a mês + YTD) =================
ppy=wb["POR PRODUTO YTD"]
def num(v):
    return round(float(v),2) if isinstance(v,(int,float)) else None
# blocos: (canal, primeira_linha_produto, ultima_linha_produto)
blocos={"Geral":(7,17),"Salão":(25,35),"Online":(43,53),"Parcerias":(61,71)}
desemp={}
for canal,(r0,r1) in blocos.items():
    arr=[]
    for r in range(r0,r1+1):
        nome=ppy.cell(r,1).value
        if not nome: continue
        nome=canon(nome)
        meses_d=[]
        for m in range(12):
            c=2+4*m
            meses_d.append([num(ppy.cell(r,c).value),num(ppy.cell(r,c+1).value),
                            ppy.cell(r,c+2).value if isinstance(ppy.cell(r,c+2).value,(int,float)) else 0,
                            None])  # ating recomputado no front
        ytd=[num(ppy.cell(r,50).value),num(ppy.cell(r,51).value),
             ppy.cell(r,52).value if isinstance(ppy.cell(r,52).value,(int,float)) else 0]
        arr.append(dict(nome=nome,meses=meses_d,ytd=ytd))
    desemp[canal]=arr
DADOS["desemp"]=desemp
json.dump(DADOS,open("dados_full.json","w",encoding="utf-8"),ensure_ascii=False,default=str)

# verificação
print("\nV6:")
for canal in ["Geral","Salão","Online","Parcerias"]:
    tot_real=sum((p["ytd"][1] or 0) for p in desemp[canal])
    tot_meta=sum((p["ytd"][0] or 0) for p in desemp[canal])
    print(f"  {canal}: {len(desemp[canal])} produtos | Real YTD {tot_real/1e6:.1f} mi | Meta YTD {tot_meta/1e6:.1f} mi")
# checar NIK FREI CANECA geral real YTD vs base(110.8)
fc=[p for p in desemp["Geral"] if p["nome"]=="NIK FREI CANECA"][0]
print("  NIK FREI CANECA Geral: Real YTD",round((fc['ytd'][1] or 0)/1e6,1),"mi | meses real:",[round((mm[1] or 0)/1e6,1) for mm in fc['meses'][:7]])


# ===== Enviar ao Firestore (roda no robô do GitHub Actions) =====
import os, json as _json, time as _time
_key=os.environ.get("FIREBASE_KEY")
if _key:
    import firebase_admin
    from firebase_admin import credentials, firestore
    cred=credentials.Certificate(_json.loads(_key))
    firebase_admin.initialize_app(cred)
    _db=firestore.client()
    _payload=_json.dumps(DADOS, ensure_ascii=False)
    _db.collection("dados").document("atual").set({
        "payloadStr": _payload,
        "atualizado": DADOS.get("atualizado"),
        "ts": int(_time.time()*1000)
    })
    print("Firestore atualizado:", len(_payload), "chars | atualizado:", DADOS.get("atualizado"))
else:
    print("FIREBASE_KEY ausente — só extraí (modo teste local).")
